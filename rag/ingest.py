"""Document ingestion: load -> clean -> chunk -> embed -> store in ChromaDB.

Supported inputs: .pdf, .docx, .md, .txt and http(s) URLs.
Ingestion is idempotent - chunk IDs are derived from (source, chunk index), so
re-running on the same file overwrites instead of duplicating.
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Iterable, List

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import Settings, settings as default_settings
from .llm import get_embeddings

SUPPORTED = {".pdf", ".docx", ".md", ".txt"}


# ----------------------------------------------------------------------------
# Loaders
# ----------------------------------------------------------------------------
def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pdf(path: Path) -> List[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = _clean(page.extract_text() or "")
        if text:
            docs.append(Document(page_content=text, metadata={"source": path.name, "page": i + 1}))
    return docs


def load_docx(path: Path) -> List[Document]:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = _clean("\n".join(parts))
    return [Document(page_content=text, metadata={"source": path.name, "page": 1})] if text else []


def load_text(path: Path) -> List[Document]:
    text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
    return [Document(page_content=text, metadata={"source": path.name, "page": 1})] if text else []


def load_url(url: str) -> List[Document]:
    import requests
    from bs4 import BeautifulSoup

    resp = requests.get(url, timeout=30, headers={"User-Agent": "Mozilla/5.0 (rag-assistant)"})
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    text = _clean(soup.get_text("\n"))
    title = soup.title.string.strip() if soup.title and soup.title.string else url
    return [Document(page_content=text, metadata={"source": url, "title": title, "page": 1})] if text else []


def load_path(path: Path) -> List[Document]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in {".md", ".txt"}:
        return load_text(path)
    raise ValueError(f"Unsupported file type: {path.name} (supported: {sorted(SUPPORTED)})")


# ----------------------------------------------------------------------------
# Chunking
# ----------------------------------------------------------------------------
def chunk_documents(docs: Iterable[Document], cfg: Settings = default_settings) -> List[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=cfg.chunk_size,
        chunk_overlap=cfg.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        add_start_index=True,
    )
    chunks = splitter.split_documents(list(docs))
    # Stable, per-source chunk numbering + deterministic id.
    counters: dict[str, int] = {}
    for c in chunks:
        src = c.metadata["source"]
        idx = counters.get(src, 0)
        counters[src] = idx + 1
        c.metadata["chunk"] = idx
        c.metadata["id"] = hashlib.sha1(f"{src}::{idx}".encode()).hexdigest()[:16]
    return chunks


# ----------------------------------------------------------------------------
# Vector store
# ----------------------------------------------------------------------------
def get_vectorstore(cfg: Settings = default_settings) -> Chroma:
    cfg.chroma_dir.mkdir(parents=True, exist_ok=True)
    return Chroma(
        collection_name=cfg.collection,
        embedding_function=get_embeddings(cfg),
        persist_directory=str(cfg.chroma_dir),
        collection_metadata={"hnsw:space": "cosine"},
    )


def index_chunks(chunks: List[Document], cfg: Settings = default_settings) -> int:
    if not chunks:
        return 0
    vs = get_vectorstore(cfg)
    ids = [c.metadata["id"] for c in chunks]
    # Chroma upserts on duplicate ids -> idempotent re-ingestion.
    batch = 64
    for i in range(0, len(chunks), batch):
        vs.add_documents(chunks[i : i + batch], ids=ids[i : i + batch])
    return len(chunks)


def ingest_paths(paths: Iterable[Path], cfg: Settings = default_settings) -> dict:
    """Ingest files. Returns {"files": n, "pages": n, "chunks": n, "skipped": [...]}."""
    all_docs: List[Document] = []
    skipped: List[str] = []
    n_files = 0
    for p in paths:
        p = Path(p)
        if p.suffix.lower() not in SUPPORTED:
            skipped.append(p.name)
            continue
        docs = load_path(p)
        if not docs:
            skipped.append(f"{p.name} (no extractable text - scanned PDF?)")
            continue
        all_docs.extend(docs)
        n_files += 1
    chunks = chunk_documents(all_docs, cfg)
    n = index_chunks(chunks, cfg)
    return {"files": n_files, "pages": len(all_docs), "chunks": n, "skipped": skipped}


def ingest_url(url: str, cfg: Settings = default_settings) -> dict:
    docs = load_url(url)
    chunks = chunk_documents(docs, cfg)
    return {"files": 1, "pages": len(docs), "chunks": index_chunks(chunks, cfg), "skipped": []}


def ingest_directory(directory: Path | None = None, cfg: Settings = default_settings) -> dict:
    directory = Path(directory or cfg.docs_dir)
    files = sorted(p for p in directory.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED)
    return ingest_paths(files, cfg)


def list_sources(cfg: Settings = default_settings) -> dict[str, int]:
    """Return {source: chunk_count} for everything in the store."""
    vs = get_vectorstore(cfg)
    data = vs.get(include=["metadatas"])
    counts: dict[str, int] = {}
    for m in data["metadatas"]:
        counts[m["source"]] = counts.get(m["source"], 0) + 1
    return dict(sorted(counts.items()))


def delete_source(source: str, cfg: Settings = default_settings) -> int:
    vs = get_vectorstore(cfg)
    data = vs.get(where={"source": source}, include=[])
    if data["ids"]:
        vs.delete(ids=data["ids"])
    return len(data["ids"])


def reset_store(cfg: Settings = default_settings) -> None:
    vs = get_vectorstore(cfg)
    vs.reset_collection()
